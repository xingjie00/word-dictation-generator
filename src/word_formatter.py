# src/word_formatter.py
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn


class WordFormatter:
    """Word文档格式化工具类"""

    @staticmethod
    def set_chinese_font(run, font_name='Microsoft YaHei'):
        """设置中文字体"""
        try:
            run.font.name = font_name
            run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
            return True
        except Exception as e:
            print(f"设置字体 {font_name} 时出错: {e}")
            try:
                run.font.name = 'SimHei'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')
                return True
            except:
                return False

    @staticmethod
    def add_text_with_font(paragraph, text, font_name='Microsoft YaHei', **kwargs):
        """添加文本并设置字体"""
        run = paragraph.add_run(text)

        if 'size' in kwargs:
            run.font.size = kwargs['size']
        if 'bold' in kwargs:
            run.font.bold = kwargs['bold']
        if 'italic' in kwargs:
            run.font.italic = kwargs['italic']
        if 'color' in kwargs:
            run.font.color.rgb = kwargs['color']

        WordFormatter.set_chinese_font(run, font_name)
        return run

    @staticmethod
    def set_cell_text_with_font(cell, text, font_name='Microsoft YaHei', **kwargs):
        """设置单元格文本并指定字体"""
        cell.text = text
        if cell.paragraphs and cell.paragraphs[0].runs:
            run = cell.paragraphs[0].runs[0]

            if 'bold' in kwargs:
                run.font.bold = kwargs['bold']
            if 'italic' in kwargs:
                run.font.italic = kwargs['italic']

            WordFormatter.set_chinese_font(run, font_name)
        return cell

    @staticmethod
    def set_global_chinese_font(doc, font_name='Microsoft YaHei'):
        """设置全局中文字体"""
        try:
            style = doc.styles['Normal']
            style.font.name = font_name
            style._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
            return True
        except Exception as e:
            print(f"设置全局字体失败: {e}")
            return False